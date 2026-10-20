"""Author the scientific manuscript only after the numerical prewriting gate."""
from pathlib import Path
import os,sys,json,re,subprocess,importlib.util,shutil
import pandas as pd
from analysis_tools import final_synthesis as s
from analysis_tools.final_documents import runtime,operation_marker,ROOT,DEPS,SKILLS

TITLE='Benchmark dependence and training adequacy shape enhancer sequence model evaluation'

def content():
    v=s.jd('claim_audit/verified_values.json');parts=[]
    def p(section,text,claims=()):parts.append(dict(section=section,text=text,claim_ids=list(claims)))
    def n(cid,key,decimals=4):return f'{v[cid][key]:.{decimals}f}'
    def interval(cid):return f'{n(cid,"estimate_a_minus_b" if "estimate_a_minus_b" in v[cid] else "estimate")} (95% CI {n(cid,"lower_95")} to {n(cid,"upper_95")})'
    p('Abstract',f'Enhancer sequence models are often compared across benchmarks that differ in sequence dependence, composition, training and experimental outcome. We audited a legacy enhancer benchmark and integrated similarity-grouped classification with quantitative reporter-activity evaluation. One hundred strong test sequences were identical to supplied training records. After fixed grouping and exclusions, the classification tasks contained 2,838 and 1,376 candidates. Current three-seed fine-tuned DNABERT-2 and DNABERT ensembles differed in pooled AUROC by {n("T_C_S1_DB2_minus_finetuned_dnabert_legacy","estimate_a_minus_b")} and {n("T_C_S2_DB2_minus_finetuned_dnabert_legacy","estimate_a_minus_b")}; conditional intervals included zero in both tasks. A separate 54-fit validation-only search compared DNABERT-2, LegNet and MPRAnn on 200 bp HepG2 and K562 reporter elements. Longer DNABERT-2 training improved internal Pearson correlation over its own two-epoch prefixes by {n("T_D_HepG2_internal_full_minus_own_early2","estimate_a_minus_b")} and {n("T_D_K562_internal_full_minus_own_early2","estimate_a_minus_b")}. LegNet nevertheless achieved higher internal correlation under the tested recipes. External DHS-natural rank correlations remained positive, while all selected-model external R² values were negative. Older frozen-model filtering and reliability diagnostics further identified high-activity underprediction. These results support careful separation of benchmark validity, training adequacy and assay transfer, rather than a universal architecture ranking. All comparisons concern previously examined cohorts and limited training searches.',['A_Stage1','A_Stage2','RUNTIME_COUNTS','T_C_S1_DB2_minus_finetuned_dnabert_legacy','T_C_S2_DB2_minus_finetuned_dnabert_legacy','T_D_HepG2_internal_full_minus_own_early2','T_D_K562_internal_full_minus_own_early2','P_MPRA'])
    p('Introduction','Enhancers contribute to cell-specific transcription, but predicting their activity from sequence remains sensitive to the outcome being measured. Reporter assays quantify activity in a specified experimental context; they do not reproduce every feature of endogenous enhancer–promoter regulation. Systematic comparisons of chromosomal and episomal assays, and of reporter design choices, show why assay context matters [@inoue_2017;@klein_2020]. Large-scale lentiMPRA measurements now provide a basis for learning quantitative activity in specific cell types [@mpra_legnet].')
    p('Introduction','Legacy enhancer benchmarks instead commonly separate enhancer from non-enhancer sequences and then strong from weak enhancer labels [@ienhancer2016]. Genomic language models, including DNABERT and DNABERT-2, learn sequence representations through pretraining and can be adapted to downstream tasks [@dnabert_original;@dnabert2]. Their evaluation requires informative conventional baselines as well as an explicit account of the training and scoring procedure. A model’s success on a binary benchmark need not translate into accurate prediction of continuous reporter activity.')
    p('Introduction','Three issues complicate comparisons. First, training and evaluation sequences may be identical or related despite different record names. Second, sequence composition can provide substantial label-associated signal, so an apparent gain may depend on the population used for evaluation. Third, unequal adaptation budgets can obscure whether a model was trained adequately. Regulatory genomics benchmarks already demonstrate the importance of strong supervised controls and task-specific evaluation [@tang2025;@dart2024]. Consequently, a further model comparison is most useful when it identifies the scope and limits of the evidence rather than reporting an isolated leaderboard.')
    p('Introduction','We therefore combine a sequence-level audit of supplied legacy files, fixed similarity-grouped classification, evaluation within exact GC counts, repeated DNABERT training, a bounded validation-only MPRA search and external assay transfer. The duration experiment compares a longer DNABERT-2 trajectory with its own preserved early prefix. Record-level prediction checks, fit failures, selection rules and distinct uncertainty summaries remain traceable. The contribution is this integrated evaluation and provenance framework; it is not a new architecture or a claim to the first language-model/CNN comparison on MPRA data. Figure 1 summarizes the design.',['RUNTIME_COUNTS'])
    p('Results|Audit reveals nonindependence in the supplied benchmark','The supplied Stage 1 benchmark contains 2,968 training and 400 test sequences, while Stage 2 contains 1,484 training and 200 test sequences. In each task, 100 strong test sequences are identical to sequences in the supplied training files. Reverse-complement matching identified no additional training–test matches in these supplied records. The exact duplicates belong to the strong class, with no duplicated weak-class test records (Table 2; Fig. 2). These are findings about the supplied files; the origin of their discrepancy with claims of an independent publisher test remains unresolved.',['A_Stage1','A_Stage2'])
    p('Results|Audit reveals nonindependence in the supplied benchmark','The optional historical three-model leakage-impact experiment could not meet its authentication requirement: the original neural historical training implementation, checkpoints and independent-test predictions are unavailable. We did not substitute an approximate modern reconstruction. A separately preserved archived spaced-mer analysis permits descriptive rescoring of the same historical predictions after duplicate test removal. That removal changes Stage 1 from 200 positives/200 negatives to 100/200, and leaves only 100 negatives in Stage 2. Cleaned Stage 2 discrimination is therefore undefined. Differences between these subsets and grouped cross-validation are not a causal estimate of leakage inflation.',['A_Stage1','A_Stage2'])
    p('Results|Grouped evaluation separates enhancer detection from strength classification',f'After historical-test priority and related-component exclusions, Stage 1 retains 2,838 candidates in 2,810 groups and Stage 2 retains 1,376 candidates in 1,373 groups. These fixed five-fold cohorts evaluate binary enhancer detection and strong-versus-weak classification, respectively (Table 1). The current DNABERT-2 three-seed ensemble attains pooled AUROC {n("C_S1_finetuned_dnabert2","auroc")} on Stage 1 and {n("C_S2_finetuned_dnabert2","auroc")} on Stage 2. The archived spaced-mer baseline attains {n("C_S1_spaced6_4_logreg","auroc")} and {n("C_S2_spaced6_4_logreg","auroc")} on the same record sets. Thus strong detection performance does not establish equally strong discrimination of the legacy strength labels.',['A_Stage1','A_Stage2','C_S1_finetuned_dnabert2','C_S2_finetuned_dnabert2','C_S1_spaced6_4_logreg','C_S2_spaced6_4_logreg'])
    p('Results|Grouped evaluation separates enhancer detection from strength classification','Table 3 keeps seven archived single-run baselines separate from the two current three-seed fine-tuned models. Mean-fold and pooled AUROC, average precision, MCC and the corresponding calibration/confusion statistics are retained. An ensemble-versus-single-run contrast describes these implementations and their aggregation choices; it does not isolate architecture or pretraining. The six mixed-label Stage 1 candidate records remain in the primary cohort, with the existing evaluation-only exclusion sensitivity reported separately (Supplementary Tables S2, S3 and S5).',['A_Stage1'])
    p('Results|Composition explains part of the apparent classification signal',f'For the archived GC/CpG logistic baseline, Stage 1 mean-fold AUROC is {n("C_S1_gc_logreg","mean_fold_auroc")}, but evaluation restricted to opposite-label pairs with identical GC count yields {n("C_S1_gc_logreg","exact_GC_mean_fold_AUROC")}. The archived spaced-mer model changes from {n("C_S1_spaced6_4_logreg","mean_fold_auroc")} to {n("C_S1_spaced6_4_logreg","exact_GC_mean_fold_AUROC")} under the same fixed rule. The current DNABERT-2 ensemble retains exact-GC pair AUROC {n("C_S1_finetuned_dnabert2","exact_GC_mean_fold_AUROC")} for Stage 1 and {n("C_S2_finetuned_dnabert2","exact_GC_mean_fold_AUROC")} for Stage 2 (Fig. 3; Supplementary Table S4). These diagnostics show that discrimination depends partly on composition and on which pairs are eligible for comparison. They do not estimate a causal contribution of GC or control every composition feature.',['C_S1_gc_logreg','C_S1_spaced6_4_logreg','C_S1_finetuned_dnabert2','C_S2_finetuned_dnabert2'])
    p('Results|Repeated training gives modest separation between DNABERT generations',f'Across the current consistent training runs, the mean of per-seed mean-fold AUROCs is {n("C_S1_finetuned_dnabert2","mean_fold_auroc")} for DNABERT-2 and {n("C_S1_finetuned_dnabert_legacy","mean_fold_auroc")} for DNABERT on Stage 1. On Stage 2, these values are {n("C_S2_finetuned_dnabert2","mean_fold_auroc")} and {n("C_S2_finetuned_dnabert_legacy","mean_fold_auroc")}. Conditional pooled-ensemble differences are {interval("T_C_S1_DB2_minus_finetuned_dnabert_legacy")} and {interval("T_C_S2_DB2_minus_finetuned_dnabert_legacy")}, respectively. Both intervals include zero, and both final-synthesis Holm-adjusted prediction-swap p-values are {n("T_C_S1_DB2_minus_finetuned_dnabert_legacy","p_holm",3)}. These results do not establish equivalence (Fig. 4).',['C_S1_finetuned_dnabert2','C_S1_finetuned_dnabert_legacy','C_S2_finetuned_dnabert2','C_S2_finetuned_dnabert_legacy','T_C_S1_DB2_minus_finetuned_dnabert_legacy','T_C_S2_DB2_minus_finetuned_dnabert_legacy'])
    p('Results|Repeated training gives modest separation between DNABERT generations','The per-seed results and descriptive fold spread quantify different forms of variation. Folds share training examples, and the effective seed schedule also reuses some seed values across different folds. The seed-by-fold entries therefore cannot be treated as independent experimental replicates. Probability reliability and fixed-threshold confusion counts are provided alongside AUROC rather than inferred from rank discrimination alone (Supplementary Figs. SF1–SF2).',['P_CLASS'])
    p('Results|Legacy enhancer labels transfer weakly to measured activity',f'In the separately archived transfer experiment, legacy Stage 1 classifiers were applied without activity-specific adaptation to top/bottom MPRA activity cohorts. DNABERT-2 score–activity Spearman correlations are {n("Z_HepG2_finetuned_dnabert2","spearman_activity")} in HepG2 and {n("Z_K562_finetuned_dnabert2","spearman_activity")} in K562. The corresponding DNABERT values are {n("Z_HepG2_finetuned_dnabert_legacy","spearman_activity")} and {n("Z_K562_finetuned_dnabert_legacy","spearman_activity")}. These are older classifier fits and different cohorts from the current repeated-seed classification and chromosome-held-out regression studies. They support a limited observation of imperfect transfer between the two outcomes; their difference from direct-regression scores cannot isolate the causal effect of changing the training task.',['Z_HepG2_finetuned_dnabert2','Z_K562_finetuned_dnabert2','Z_HepG2_finetuned_dnabert_legacy','Z_K562_finetuned_dnabert_legacy'])
    p('Results|Direct MPRA training recovers substantial sequence signal',f'Direct activity modelling uses 200 bp elements and frozen chromosome partitions, with 9,048 internal test elements in HepG2 and 13,145 in K562. The selected DNABERT-2 ensembles achieve Pearson correlations {n("M_HepG2_dnabert2_selected","pearson")} and {n("M_K562_dnabert2_selected","pearson")}. LegNet achieves {n("M_HepG2_legnet_selected","pearson")} and {n("M_K562_legnet_selected","pearson")}, while MPRAnn achieves {n("M_HepG2_mprann_selected","pearson")} and {n("M_K562_mprann_selected","pearson")}. Table 4 reports Spearman, RMSE, MAE and R² as well as Pearson. The archived spaced-mer regression baseline remains a separate single recipe outside the neural search. These results demonstrate useful sequence-associated signal for the tested reporter tasks, without establishing a universal limit or advantage for a model family (Fig. 5).',['D_HepG2','D_K562','P_MPRA','M_HepG2_dnabert2_selected','M_K562_dnabert2_selected','M_HepG2_legnet_selected','M_K562_legnet_selected','M_HepG2_mprann_selected','M_K562_mprann_selected'])
    p('Results|Longer DNABERT training helps while LegNet remains stronger',f'The fixed-learning-rate duration control compares each longer DNABERT-2 trajectory with its own exact first-two-epoch prefix. Internal Pearson improves by {interval("T_D_HepG2_internal_full_minus_own_early2")} in HepG2 and {interval("T_D_K562_internal_full_minus_own_early2")} in K562. Both Holm-adjusted conditional prediction-swap p-values are {n("T_D_HepG2_internal_full_minus_own_early2","p_holm")}. This addresses the concern that the earlier short training budget limited performance, using a duration contrast separated from validation-selected model comparisons.',['T_D_HepG2_internal_full_minus_own_early2','T_D_K562_internal_full_minus_own_early2'])
    p('Results|Longer DNABERT training helps while LegNet remains stronger',f'After validation-only selection, LegNet minus DNABERT-2 Pearson differences are {interval("T_M_HepG2_internal_legnet_selected_minus_dnabert2_selected")} and {interval("T_M_K562_internal_legnet_selected_minus_dnabert2_selected")} for HepG2 and K562, with adjusted conditional p-values {n("T_M_HepG2_internal_legnet_selected_minus_dnabert2_selected","p_holm")}. DNABERT-2 versus MPRAnn is less clear: the HepG2 interval includes zero, while the positive K562 pointwise interval accompanies Holm-adjusted p={n("T_M_K562_internal_dnabert2_selected_minus_mprann_selected","p_holm")}. Pointwise intervals must therefore be interpreted separately from the adjusted test family (Table 6).',['T_M_HepG2_internal_legnet_selected_minus_dnabert2_selected','T_M_K562_internal_legnet_selected_minus_dnabert2_selected','T_M_HepG2_internal_dnabert2_selected_minus_mprann_selected','T_M_K562_internal_dnabert2_selected_minus_mprann_selected'])
    p('Results|Longer DNABERT training helps while LegNet remains stronger',f'Training adequacy remains bounded. Six of the 18 selected fits are potentially budget-limited, and the LegNet and MPRAnn choices lie on opposite grid boundaries. Two high-learning-rate DNABERT-2 fits are practically near constant despite finite validation Pearson; zero fits meet the original undefined-Pearson flag. All fits remain in the original three-seed candidate means. Across nine candidate/seed fits per model and cell, recorded DNABERT-2 runtime totals {n("K_HepG2_dnabert2","elapsed_hours",2)} and {n("K_K562_dnabert2","elapsed_hours",2)} hours, compared with {n("K_HepG2_legnet","elapsed_hours",3)} and {n("K_K562_legnet","elapsed_hours",3)} for LegNet. These are measured fitting times, not standardized FLOPs or total pretraining costs (Supplementary Tables S7, S16–S17).',['RUNTIME_COUNTS','K_HepG2_dnabert2','K_K562_dnabert2','K_HepG2_legnet','K_K562_legnet'])
    p('Results|External evaluation preserves ranking more than absolute calibration',f'On the frozen DHS-natural cohort of 8,399 sequences per cell, DNABERT-2 Spearman correlation is {n("E_HepG2_dnabert2_selected","spearman")} in HepG2 and {n("E_K562_dnabert2_selected","spearman")} in K562. LegNet reaches {n("E_HepG2_legnet_selected","spearman")} and {n("E_K562_legnet_selected","spearman")}, while MPRAnn reaches {n("E_HepG2_mprann_selected","spearman")} and {n("E_K562_mprann_selected","spearman")}. LegNet has positive conditional paired differences from DNABERT-2 in both cells under this selected-model comparison (Table 6; Fig. 6A).',['E_HepG2_dnabert2_selected','E_K562_dnabert2_selected','E_HepG2_legnet_selected','E_K562_legnet_selected','E_HepG2_mprann_selected','E_K562_mprann_selected','T_M_HepG2_DHS_natural_legnet_selected_minus_dnabert2_selected','T_M_K562_DHS_natural_legnet_selected_minus_dnabert2_selected'])
    p('Results|External evaluation preserves ranking more than absolute calibration','All six selected-model external R² values are negative, and all mean prediction-minus-measurement errors are negative. Longer DNABERT-2 training improves external rank correlation but worsens absolute prediction error in both cells, as retained in the fixed-prefix secondary diagnostics. Thus rank preservation does not demonstrate accurate transfer of absolute activity units. This cohort was previously examined; it cannot provide a new confirmatory holdout or guide further model selection (Table 5; Fig. 6B; Supplementary Table S8).',['E_HepG2_dnabert2_selected','E_K562_dnabert2_selected','E_HepG2_legnet_selected','E_K562_legnet_selected','E_HepG2_mprann_selected','E_K562_mprann_selected'])
    p('Results|Filtering and activity tails reveal different failure modes',f'The complete paired-assay pool contains 149,632 candidates across the two cells, of which 76,833 are retained and 72,799 are removed by the original overlap screen. Most removals are triggered by shared short words, rather than full-sequence duplication. Older frozen-model transport analyses compare retained and removed test-chromosome candidates. For the older DNABERT-2 fits, removed-minus-retained MAE is {interval("F_HepG2_dnabert2")} in HepG2 and {interval("F_K562_dnabert2")} in K562. Both exploratory intervals contain zero; these comparisons do not establish the benefit of retraining without filtering (Supplementary Tables S10–S12).',['RUNTIME_COUNTS','F_HepG2_dnabert2','F_K562_dnabert2'])
    p('Results|Filtering and activity tails reveal different failure modes',f'The older CNN diagnostics show persistent underprediction among retained high-activity elements. For example, the mean signed errors of the older LegNet fits are {n("H_HepG2_legnet","signed_error")} for HepG2 and {n("H_K562_legnet","signed_error")} for K562. Across-seed agreement and marginal interval coverage do not guarantee accurate predictions in this tail. These analyses use older frozen models, fixed training-derived thresholds and measured-activity bands; they do not evaluate the newly selected models or define a prospective target-free selection rule (Fig. 6C–D; Supplementary Tables S13–S14).',['H_HepG2_legnet','H_K562_legnet'])
    p('Discussion','The benchmark audit changes the interpretation of apparently independent evaluation. Different record names are insufficient evidence of independence when sequences are identical or connected by a stated similarity rule. The observed overlap is directly verifiable in the supplied files. Its causal effect on unavailable original neural historical results is not identifiable here, and the cleaned strength test becomes degenerate. Reporting this limit is more informative than assigning an approximate inflation estimate to an unauthenticated historical implementation. Grouping also does not rule out every form of homology, shared lineage or pretraining membership.')
    p('Discussion','Model advantage is task- and recipe-dependent. The repeated classification study does not show a resolved difference between DNABERT generations on these fixed cohorts. In quantitative activity prediction, extra DNABERT-2 training materially improves the tested short-prefix result, yet LegNet remains stronger on the selected internal Pearson comparison. This is consistent with existing evidence that specialized supervised controls can be competitive in regulatory genomics [@tang2025;@dart2024]. Our finite candidate search narrows an obvious budget concern without establishing equal-compute fairness or an exhaustive architecture optimum.')
    p('Discussion','Legacy enhancer labels and reporter activity should remain separate scientific targets. The former summarize a curated classification problem; the latter reflect an assay, cellular setting and construct design. Imperfect score transfer between them does not prove that sequence contains no activity information. The direct regression results instead show useful predictive information within the tested assay. Conversely, activity in a reporter construct should not be interpreted as complete endogenous regulatory function, which also depends on cellular and genomic context [@inoue_2017;@klein_2020;@fulco_2019].')
    p('Discussion','External evaluation highlights the separation between ranking and calibration. A positive rank correlation may still accompany predictions that perform worse in squared error than an external-cohort mean predictor. The current negative external R² values and signed bias are therefore substantive findings, rather than secondary details to omit from a correlation leaderboard. The older activity-tail and coverage diagnostics further show why ensemble agreement alone is inadequate evidence of individual reliability. Marginal calibration does not guarantee conditional accuracy across all subpopulations [@barber2021;@barber2023].')
    p('Discussion','Several limitations bound the conclusions. The activity inputs are short sequences from two cell types, and the historical classification outcomes are a separate task. Only three training seeds are available for the current neural studies. Fixed folds share training data, and the conditional cohort analyses exclude uncertainty from model selection, retraining and hyperparameter search. The inferential family was chosen for final synthesis after prior outcomes had been seen; multiplicity adjustment does not convert that history into prospective confirmation. Reusing examined cohorts warrants caution in the interpretation of new analyses [@bates2024;@dwork2015]. Practical collapse, selected budget caps, boundary choices and the unresolved earlier locus-split collapse prevent broad convergence claims. Biological mechanism, longer context, other cell types and universal model superiority remain outside this study.',['P_CLASS','P_MPRA','RUNTIME_COUNTS'])
    p('Methods|Source authentication and benchmark design','Source archives, manifests, preprocessing configurations, prediction files and selected checkpoints were tracked by SHA256. Multiline FASTA sequences were validated and exact/reverse-complement identities audited. The legacy graph connects sequences at longest common subsequence length at least 160 in either orientation and uses connected components for the fixed five-fold split. This is a sequence grouping rule, not an 80% local identity threshold or chromosome-disjoint evaluation. Test-containing components are excluded from candidate training. Exact-GC diagnostics compare every eligible opposite-label pair within the same held-out fold, without retraining. Supplementary Methods provide the parsing and grouping pseudocode, all exclusions and the matching estimand.',['A_Stage1','A_Stage2'])
    p('Methods|Classification training and scoring','The current repeated comparison uses DNABERT and DNABERT-2 at pinned revisions with attention-mask mean pooling, dropout and a two-class linear head [@dnabert_original;@dnabert2]. Base seeds are 20260928, 20261016 and 20261017. Each model, task, seed and fixed fold is trained for two epochs using AdamW at learning rate 2e-5, batch size 8 and gradient accumulation 2. Reverse complements augment training after splitting, while scoring uses the original orientation. The fixed final second epoch and probability threshold 0.5 are used; outer-fold outcomes select neither epochs nor thresholds. The seven remaining baselines retain their original authenticated one-run recipes, including frozen GENA-LM [@gena2025]. Full settings and the effective seed schedule accompany the supplement.',['P_CLASS'])
    p('Methods|Activity models and validation selection','The current activity study fits DNABERT-2, the copied human MPRA LegNet architecture and MPRAnn implementations [@mpra_legnet;@penzar2023legnet]. Each model/cell receives three declared learning-rate candidates and three seeds (20260930–20260932), giving 54 fits. Train-only target mean and population SD are used for normalization. DNABERT-2 permits up to 30 epochs with patience 5 and minimum 8; LegNet retains a 20-epoch OneCycle schedule; MPRAnn permits 100 epochs with patience 10. Checkpoints use the highest finite validation Pearson with earliest-epoch ties. Recipe selection uses the mean validation Pearson across all three seeds; an undefined seed makes a candidate ineligible. All recipe/checkpoint hashes are locked before new test/external inference. Full grids, model-specific batching, augmentation, precision and stopping details are specified in Supplementary Methods.',['P_MPRA','RUNTIME_COUNTS'])
    p('Methods|Duration controls and external evaluation','Every DNABERT-2 fit preserves the validation-best checkpoint within its first two epochs. The primary duration comparison fixes learning rate 2e-5 and compares each full fit with its own identical early trajectory. Activity predictions average forward and reverse-complement scores, then ensemble the three seeds. Internal evaluation uses the frozen primary chromosome test. Current external evaluation uses the frozen DHS-natural cohort; other origin and filtering analyses are explicitly labelled by their older fitted-model generation. Spearman is the external headline, with Pearson, MAE, RMSE, R² and signed error reported to retain calibration failures. No external or test outcome changes the selection policy.',['P_MPRA','D_HepG2','D_K562'])
    p('Methods|Statistical synthesis and reproducibility','The final synthesis family contains 14 headline contrasts fixed after earlier outcomes were examined. Conditional percentile intervals use 2,000 paired whole-group bootstrap draws: similarity-graph components for classification, locus groups internally and canonical-DNA groups externally. Spearman ranks are recomputed for duplicated sampled observations. Two-sided paired prediction swaps use 9,999 groupwise permutations and a plus-one Monte Carlo correction. Their sharp conditional exchangeability null differs from algorithm equality or equivalence. Holm adjustment controls the stated family of conditional tests [@holm1979]; pointwise intervals, Monte Carlo precision and seed/fold variability remain distinct. The original external 1 kb region sensitivity and older unadjusted secondary diagnostics remain separately labelled. The checkpoint audit reproduces 60 classification checkpoints and 72 MPRA validation checkpoints plus 60 internal/external evaluation sets. Reporting maps follow DOME and REFORMS topics without asserting certification or complete elimination of leakage [@dome2021;@reforms2024].',['P_STATS','RUNTIME_COUNTS'])
    return parts

def references(parts):
    source=s.read('phase27_outputs/verified_primary_citations.csv').fillna('').set_index('source_id').to_dict('index')
    additions={
      'ienhancer2016':dict(authors='Liu, B., Fang, L., Long, R., Lan, X. and Chou, K.-C.',year=2016,title='iEnhancer-2L: a two-layer predictor for identifying enhancers and their strength by pseudo k-tuple nucleotide composition',doi='10.1093/bioinformatics/btv604',url='https://academic.oup.com/bioinformatics/article/32/3/362/1744331',publication_status='peer-reviewed'),
      'dome2021':dict(authors='Walsh, I. et al.',year=2021,title='DOME: recommendations for supervised machine learning validation in biology',doi='10.1038/s41592-021-01205-4',url='https://www.nature.com/articles/s41592-021-01205-4',publication_status='peer-reviewed recommendation'),
      'reforms2024':dict(authors='Kapoor, S. et al.',year=2024,title='REFORMS: Consensus-based Recommendations for Machine-learning-based Science',doi='10.1126/sciadv.adk3452',url='https://pubmed.ncbi.nlm.nih.gov/38691601/',publication_status='peer-reviewed recommendation'),
      'holm1979':dict(authors='Holm, S.',year=1979,title='A simple sequentially rejective multiple test procedure',doi='10.2307/4615733',url='https://www.jstor.org/stable/4615733',publication_status='peer-reviewed')}
    source.update(additions);order=[]
    for part in parts:
        for group in re.findall(r'\[@([^]]+)\]',part['text']):
            for sid in group.split(';@'):
                assert sid in source,sid
                if sid not in order:order.append(sid)
    refs=[dict(reference_number=i+1,source_id=sid,**source[sid]) for i,sid in enumerate(order)]
    def resolve(text):
        return re.sub(r'\[@([^]]+)\]',lambda m:'['+', '.join(str(order.index(sid)+1) for sid in m.group(1).split(';@'))+']',text)
    for part in parts:part['text']=resolve(part['text'])
    return refs

def word_document(parts,refs,path):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.section import WD_SECTION_START,WD_ORIENT
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(11 if name=='Normal' else 20 if name=='Title' else 14 if name=='Heading 1' else 12 if name=='Heading 2' else 9.5)
        style.paragraph_format.space_after=Pt(7)
        if name.startswith('Heading'):style.paragraph_format.keep_with_next=True
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    for sec in doc.sections:
        footer=sec.footer.paragraphs[0];footer.alignment=2
        f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    doc.add_paragraph(TITLE,'Title');current=None
    for part in parts:
        section=part['section'];main,_,sub=section.partition('|')
        if current!=section:
            if current is None or main!=(current.partition('|')[0]):doc.add_heading(main,level=1)
            if sub:doc.add_heading(sub,level=2)
            current=section
        doc.add_paragraph(part['text'])
    doc.add_heading('References',level=1)
    for r in refs:
        authors=str(r['authors']).replace(' and ', ', ')
        entry=f'{r["reference_number"]}. {authors} ({int(r["year"])}). {r["title"]}.'
        if r.get('doi'):entry+=' https://doi.org/'+r['doi']
        else:entry+=' '+r['url']
        p=doc.add_paragraph(entry);p.paragraph_format.space_after=Pt(5)
        for run in p.runs:run.font.size=Pt(9.5)
    # Main tables are presented as compact reader views of full traceable CSVs.
    def section_landscape():
        sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
        sec.page_width=Inches(11);sec.page_height=Inches(8.5);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.65)
        return sec
    def add_table(frame):
        t=doc.add_table(rows=1,cols=len(frame.columns));t.autofit=False
        available=9.7;first=2.0 if 'Model' in frame.columns else 1.45;other=(available-first)/(len(frame.columns)-1)
        widths=[first]+[other]*(len(frame.columns)-1)
        for j,name in enumerate(frame.columns):t.rows[0].cells[j].text=str(name)
        for _,r in frame.iterrows():
            cells=t.add_row().cells
            for j,value in enumerate(r):
                cells[j].text='—' if pd.isna(value) else str(value)
        for i,row in enumerate(t.rows):
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);tcpr=cell._tc.get_or_add_tcPr()
                margin=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    node=OxmlElement('w:'+edge);node.set(qn('w:w'),'85');node.set(qn('w:type'),'dxa');margin.append(node)
                tcpr.append(margin)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(3);paragraph.paragraph_format.line_spacing=1.0
                    for run in paragraph.runs:run.font.size=Pt(10);run.bold=i==0
            if i==0:
                header=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(header)
    tables=s.read('final_tables/table_index.csv');fmt=lambda x:f'{x:.4f}'
    for r in tables[tables.placement.eq('main')].itertuples():
        section_landscape();doc.add_heading('Table '+r.table.split('_')[0][1:]+' '+r.title,level=1);d=s.read(r.file)
        if r.table.startswith('T1'):
            legacy=d[d.task.str.startswith('Legacy')]
            add_table(pd.DataFrame({'Task':legacy.task,'Supplied train':legacy.initial_training.astype(int),'Historical test':legacy.historical_test.astype(int),'Grouped candidates':legacy.retained.astype(int),'Groups':legacy.groups.astype(int),'Folds':legacy.folds.astype(int)}))
            doc.add_paragraph('Reporter activity populations')
            a=d[d.task.isin(['HepG2','K562'])]
            add_table(pd.DataFrame({'Cell':a.task,'Complete pairs':a.complete_assay_candidates.astype(int),'Retained':a.retained.astype(int),'Train':a.train.astype(int),'Validation':a.validation.astype(int),'Internal test':a.internal_test.astype(int),'External':a.external_DHS_natural.astype(int)}))
        elif r.table.startswith('T2'):
            add_table(d[['stage','supplied_train_n','supplied_test_n','exact_train_test_overlap','reverse_complement_train_test_overlap','duplicated_strong_records','duplicated_weak_records']].rename(columns={'stage':'Stage','supplied_train_n':'Train','supplied_test_n':'Test','exact_train_test_overlap':'Exact overlaps','reverse_complement_train_test_overlap':'RC overlaps','duplicated_strong_records':'Strong duplicates','duplicated_weak_records':'Weak duplicates'}))
            doc.add_paragraph('Retained candidates after historical-test priority and related-group exclusions')
            add_table(d[['stage','additional_test_related_training_excluded','candidate_n','candidate_groups','mixed_label_candidate_groups','mixed_label_candidate_records']].rename(columns={'stage':'Stage','additional_test_related_training_excluded':'Additional related exclusions','candidate_n':'Candidates','candidate_groups':'Groups','mixed_label_candidate_groups':'Mixed-label groups','mixed_label_candidate_records':'Mixed-label records'}))
        elif r.table.startswith('T3'):
            for stage in [1,2]:
                doc.add_paragraph(f'Stage {stage}');a=d[d.stage.eq(stage)]
                add_table(pd.DataFrame({'Model':a.model_label,'Seeds':a.n_seeds,'Mean-fold AUROC':a.mean_fold_auroc.map(fmt),'Fold SD':a.fold_sd_of_seed_mean_auroc.map(fmt),'Seed SD':a.seed_sd_of_mean_fold_auroc.map(lambda x:'—' if pd.isna(x) else fmt(x)),'Pooled AUROC':a.auroc.map(fmt),'AUPRC':a.auprc.map(fmt),'Exact-GC AUROC':a.exact_GC_mean_fold_AUROC.map(fmt)}))
        elif r.table.startswith('T4') or r.table.startswith('T5'):
            keys=['cell','model_label','pearson','spearman','rmse','mae','r2'] if r.table.startswith('T4') else ['cell','model_label','spearman','pearson','mae','rmse','r2']
            a=d[keys].copy()
            for col in keys[2:]:a[col]=a[col].map(fmt)
            add_table(a.rename(columns={'cell':'Cell','model_label':'Model','pearson':'Pearson','spearman':'Spearman','rmse':'RMSE','mae':'MAE','r2':'R²'}))
        else:
            a=pd.DataFrame({'Task and cohort':d.apply(lambda x:f'Stage {int(x.stage)}' if x.task=='classification' else x.cell+' '+x.cohort,axis=1),'Contrast':d.contrast_label,'Difference':d.estimate_a_minus_b.map(fmt),'95% interval':d.apply(lambda x:f'{x.lower_95:.4f} to {x.upper_95:.4f}',axis=1),'Holm p':d.p_holm.map(fmt)})
            add_table(a)
        doc.add_paragraph(r.notes,'Caption')
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.PORTRAIT
    sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.left_margin=sec.right_margin=Inches(.85);sec.top_margin=sec.bottom_margin=Inches(.8)
    doc.add_heading('Figure legends',level=1)
    for r in s.read('final_figures/figure_index.csv').query('placement=="main"').itertuples():
        doc.add_heading('Figure '+r.figure.split('_')[0][1:]+' '+r.title,level=2)
        doc.add_paragraph(r.caption)
    doc.core_properties.title=TITLE;doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(ROOT/path)

def render_docx(path,outdir):
    runtime();out=ROOT/outdir;out.mkdir(parents=True,exist_ok=True)
    renderer_path=SKILLS/'documents/26.1011.11113/skills/documents/render_docx.py'
    spec=importlib.util.spec_from_file_location('study_bundled_renderer',renderer_path);renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)
    os.environ['PATH']=str(DEPS/'native/poppler/Library/bin')+os.pathsep+os.environ['PATH']
    log=[]
    try:
        pages=renderer.rasterize(str(ROOT/path),str(out),130,False,True)
        backend='packaged LibreOffice conversion'
    except FileNotFoundError as exc:
        if 'soffice' not in str(exc).lower():raise
        log.append(str(exc));pdf=out/(Path(path).stem+'.pdf')
        command=['powershell.exe','-NoProfile','-NonInteractive','-File',str(ROOT/'analysis_tools/export_final_docx_word.ps1'),'-InputPath',str(ROOT/path),'-OutputPath',str(pdf)]
        result=subprocess.run(command,cwd=ROOT,env=os.environ.copy(),capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=180,creationflags=subprocess.CREATE_NO_WINDOW)
        log.extend([result.stdout,result.stderr]);assert result.returncode==0,'Word export failed: '+result.stderr
        assert pdf.exists() and pdf.stat().st_size>0
        # Reuse the packaged renderer's rasterization and naming pipeline with
        # the native Word-exported PDF; do not alter any managed skill files.
        renderer.convert_to_pdf=lambda *args,**kwargs:(str(pdf),'Native Microsoft Word PDF export; LibreOffice unavailable on this Windows host.')
        pages=renderer.rasterize(str(ROOT/path),str(out),130,False,False)
        backend='native Word PDF export + packaged render_docx.rasterize / bundled Poppler'
    s.dump(out/'render_backend.json',dict(backend=backend,renderer_sha256=s.sha(renderer_path),document_sha256=s.sha(ROOT/path),pages=len(pages),log=log,system_settings_changed=False))
    return len(pages)

def phase46():
    runtime();s.verify_inputs();assert s.jd('phase45_outputs/summary.json')['status']=='completed_prewrite_numeric_gate_passed'
    for r in s.jd('claim_audit/prewriting_evidence_lock.json'):assert s.sha(ROOT/r['path'])==r['sha256'],r['path']
    parts=content();refs=references(parts)
    ledger=s.read('claim_audit/claim_evidence_ledger.csv').set_index('claim_id')
    for part in parts:
        for cid in part['claim_ids']:assert ledger.loc[cid,'status']=='VERIFIED',cid
    s.dump('manuscript/manuscript_paragraphs.json',parts);s.table('manuscript/references.csv',pd.DataFrame(refs))
    md='# '+TITLE+'\n\n';current=None
    for part in parts:
        if part['section']!=current:
            main,_,sub=part['section'].partition('|')
            if current is None or current.partition('|')[0]!=main:md+='## '+main+'\n\n'
            if sub:md+='### '+sub+'\n\n'
            current=part['section']
        md+=part['text']+'\n\n'
    md+='## References\n\n'
    for r in refs:md+=f'{r["reference_number"]}. {r["authors"]} ({int(r["year"])}). {r["title"]}. '+('https://doi.org/'+r['doi'] if r.get('doi') else r['url'])+'\n\n'
    s.text('manuscript/Manuscript.md',md)
    operation_marker('create',1,'docx','phase46_outputs/docx_operation_marker.json')
    word_document(parts,refs,'manuscript/Manuscript.docx')
    pages=render_docx('manuscript/Manuscript.docx','phase46_outputs/rendered_manuscript')
    docx=ROOT/'manuscript/Manuscript.docx'
    import zipfile,xml.etree.ElementTree as E
    with zipfile.ZipFile(docx) as z:
        assert z.testzip() is None
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        document=E.fromstring(z.read('word/document.xml'));actual='\n'.join(''.join(p.itertext()) for p in document.findall('.//w:p',ns))
    checks=[]
    for i,part in enumerate(parts,1):
        assert part['text'] in actual,part['section']
        checks.append(dict(paragraph=i,section=part['section'],claim_ids=';'.join(part['claim_ids']),all_referenced_numerical_claims_verified=True,exact_paragraph_in_DOCX=True))
    assert len(set(re.findall(r'\[(\d+)\]',actual)))<=len(refs)
    assert not re.search(r'\[@|TODO|TBD|INSERT HERE|citation needed',actual,re.I)
    s.table('claim_audit/actual_manuscript_paragraph_checks.csv',pd.DataFrame(checks))
    s.dump('phase46_outputs/summary.json',dict(status='drafted_and_rendered_pending_visual_QA',paragraphs=len(parts),references=len(refs),rendered_pages=pages,all_registered_numerical_claims_verified=True,all_paragraphs_present_in_docx=True,new_training=False,journal_target='not selected',author_declarations='pending specific author confirmation',word_count_body=sum(len(p['text'].split()) for p in parts),source_sha256=s.sha(__file__)))
    print('Phase46 manuscript authored and rendered:',pages,'pages; visual QA pending.')
