"""Local, reviewable release and submission artifacts; never publishes or submits."""
from pathlib import Path
import ast,hashlib,json,os,shutil,subprocess,sys,zipfile,importlib.metadata as metadata
import numpy as np
import pandas as pd
from analysis_tools import final_synthesis as s
ROOT=s.ROOT
RELEASE=ROOT/'public_release/enhancer-leakage-study'

def copy_file(relative,destination):
    source=ROOT/relative;target=Path(destination)
    target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(source,target)
    assert s.sha(source)==s.sha(target)

def write(path,text):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(text.strip()+'\n',encoding='utf-8')

def manifest(folder):
    return [dict(path=p.relative_to(folder).as_posix(),bytes=p.stat().st_size,sha256=s.sha(p)) for p in sorted(folder.rglob('*')) if p.is_file() and p.name!='MANIFEST.json' and 'reproduced' not in p.relative_to(folder).parts and '__pycache__' not in p.parts and '.runtime' not in p.parts]

def seal(folder,path):
    items=manifest(folder);write(folder/'MANIFEST.json',json.dumps(items,indent=2))
    needed=sum(r['bytes'] for r in items)
    assert shutil.disk_usage(ROOT).free>needed*2+500_000_000,'Preserve prior packages; insufficient space for a new archive.'
    temp=path.with_suffix('.zip.writing')
    with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED,compresslevel=6,allowZip64=True) as z:
        for r in items:z.write(folder/r['path'],folder.name+'/'+r['path'])
        z.write(folder/'MANIFEST.json',folder.name+'/MANIFEST.json')
    with zipfile.ZipFile(temp) as z:
        assert z.testzip() is None
        for r in items:assert hashlib.sha256(z.read(folder.name+'/'+r['path'])).hexdigest()==r['sha256']
    os.replace(temp,path)
    return dict(path=path.relative_to(ROOT).as_posix(),bytes=path.stat().st_size,sha256=s.sha(path),members=len(items)+1,all_CRCs_passed=True,all_manifest_member_hashes_passed=True)

def finish46():
    from analysis_tools.final_documents import runtime
    runtime();from pypdf import PdfReader,PdfWriter
    summary=s.jd('phase46_outputs/summary.json');assert summary['rendered_pages']==16
    qa=dict(passed=True,pages_inspected=16,method='All manuscript text, references, table and legend pages viewed; table grid widths and pagination corrected and altered pages re-inspected. Six main scientific figures previously inspected enlarged.',checks=['counts and table values','table column spacing','references and legends','no clipped text or orphan table footnote','current and archived evidence distinguished'],document_sha256=s.sha(ROOT/'manuscript/Manuscript.docx'),rendered_pdf_sha256=s.sha(ROOT/'phase46_outputs/rendered_manuscript/Manuscript.pdf'),page_png_sha256={str(i):s.sha(ROOT/f'phase46_outputs/rendered_manuscript/page-{i}.png') for i in range(1,17)})
    s.dump('phase46_outputs/visual_QA.json',qa)
    writer=PdfWriter();writer.append(ROOT/'phase46_outputs/rendered_manuscript/Manuscript.pdf')
    for r in s.read('final_figures/figure_index.csv').query('placement=="main"').itertuples():writer.append(ROOT/r.pdf)
    with (ROOT/'manuscript/Manuscript.pdf').open('wb') as f:writer.write(f)
    assert len(PdfReader(ROOT/'manuscript/Manuscript.pdf').pages)==22
    summary.update(status='completed_scientific_draft',visual_QA_passed=True,review_pdf_pages=22,submission_author_and_journal_approval_pending=True)
    s.dump('phase46_outputs/summary.json',summary)

ADAPTER='''from pathlib import Path
import hashlib,json,os
import pandas as pd
PACKAGE=Path(__file__).resolve().parents[1]
ROOT=PACKAGE/'reproduced'
def sha(path):
    with Path(path).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def read(path):return pd.read_csv(PACKAGE/'inputs'/path,float_precision='round_trip')
def jd(path):return json.loads((PACKAGE/'inputs'/path).read_text(encoding='utf-8'))
def dump(path,obj):
    p=ROOT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,indent=2,allow_nan=False)+'\\n',encoding='utf-8')
def text(path,value):
    p=ROOT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(value.strip()+'\\n',encoding='utf-8')
def table(path,frame):
    p=ROOT/path;p.parent.mkdir(parents=True,exist_ok=True);frame.to_csv(p,index=False,float_format='%.17g')
def manifests():return [read(f'final_evidence/inputs/stage{stage}_candidate_manifest.csv') for stage in [1,2]]
def verify_inputs():
    for r in json.loads((PACKAGE/'MANIFEST.json').read_text(encoding='utf-8')):assert sha(PACKAGE/r['path'])==r['sha256'],r['path']
    return True
'''

METRICS='''import sys,json,hashlib
from pathlib import Path
import numpy as np
import pandas as pd
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from analysis_tools import final_synthesis as s
from analysis_tools.metric_functions import *
s.verify_inputs();c=s.read('final_statistics/classification_predictions.csv.gz');m=s.read('phase37_outputs/verified_ensemble_predictions.csv.gz')
expected=s.read('final_statistics/classification_summary.csv').set_index(['stage','model']);rows=[];errors=[]
for (stage,model),g in c.groupby(['stage','model']):
    got=classification_metrics(g.label,g.score)
    for key,value in got.items():errors.append(abs(value-expected.loc[(stage,model),key]))
    rows.append(dict(stage=stage,model=model,**got))
s.table('classification_metrics_recomputed.csv',pd.DataFrame(rows))
expected=pd.concat([s.read('final_statistics/mpra_internal_summary.csv'),s.read('final_statistics/external_transfer_summary.csv')]).set_index(['cell','cohort','arm']);rows=[]
for (cell,cohort,arm),g in m.groupby(['cell','cohort','arm']):
    got=regression_metrics(g.activity,g.prediction)
    for key,value in got.items():errors.append(abs(value-expected.loc[(cell,cohort,arm),key]))
    rows.append(dict(cell=cell,cohort=cohort,arm=arm,**got))
s.table('MPRA_metrics_recomputed.csv',pd.DataFrame(rows));assert max(errors)<1e-12
checks=stats_selfcheck();s.table('weighted_vs_literal.csv',checks)
contrasts=s.read('final_statistics/adjusted_tests.csv');point_errors=[];stat_errors=[];results=[]
policy=s.jd('final_statistics/statistical_methods.json')
for row in contrasts.to_dict('records'):
    y,a,b,gi,ng=paired_arrays(row,c,m);w0=np.ones(len(y))
    obs=weighted_metric(y,a,w0,row['metric'])-weighted_metric(y,b,w0,row['metric'])
    point_errors.append(abs(obs-row['estimate_a_minus_b']))
    if '--full-statistics' not in sys.argv:continue
    seed=int.from_bytes(hashlib.sha256(row['comparison_id'].encode()).digest()[:4],'little')^policy['random_seed'];rng=np.random.default_rng(seed);boot=[]
    for i in range(policy['bootstrap_draws']):
        w=np.bincount(rng.integers(0,ng,ng),minlength=ng)[gi].astype(float)
        boot.append(weighted_metric(y,a,w,row['metric'])-weighted_metric(y,b,w,row['metric']))
    defined=np.asarray(boot);defined=defined[np.isfinite(defined)];lo,hi=np.quantile(defined,[.025,.975]);extreme=0
    for i in range(policy['permutation_draws']):
        flip=rng.integers(0,2,ng,dtype=np.int8)[gi].astype(bool);pa=np.where(flip,b,a);pb=np.where(flip,a,b)
        value=weighted_metric(y,pa,w0,row['metric'])-weighted_metric(y,pb,w0,row['metric']);assert np.isfinite(value)
        extreme+=abs(value)>=abs(obs)-1e-14
    p=(extreme+1)/(policy['permutation_draws']+1)
    stat_errors.extend([abs(lo-row['lower_95']),abs(hi-row['upper_95']),abs(p-row['p_raw'])]);results.append(dict(comparison_id=row['comparison_id'],estimate=obs,lower_95=lo,upper_95=hi,p_raw=p))
    print('Reproduced',len(results),'/14:',row['comparison_id'],flush=True)
assert max(point_errors)<1e-12
if results:
    f=pd.DataFrame(results);order=np.argsort(f.p_raw.to_numpy(),kind='stable');p=f.p_raw.to_numpy()[order];adj=np.minimum(1,np.maximum.accumulate(p*(len(p)-np.arange(len(p)))));out=np.empty(len(p));out[order]=adj;f['p_holm']=out
    stat_errors.extend(abs(f.p_holm.to_numpy()-contrasts.p_holm.to_numpy()));assert max(stat_errors)<1e-12;s.table('headline_statistics_recomputed.csv',f)
s.dump('metric_validation.json',dict(passed=True,classification_rows=18,MPRA_ensemble_rows=20,weighted_literal_checks=len(checks),point_contrasts=14,max_metric_error=float(max(errors)),max_point_difference_error=float(max(point_errors)),full_statistics_contrasts=len(results),max_interval_pvalue_error=float(max(stat_errors)) if stat_errors else None,bootstrap_draws=policy['bootstrap_draws'] if results else 0,permutation_draws=policy['permutation_draws'] if results else 0,scope=policy['scope']))
print('Metric reproduction passed.',flush=True)
'''

AUDIT='''import sys
from pathlib import Path
import pandas as pd
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from analysis_tools import final_synthesis as s
s.verify_inputs();roster=s.read('phase21_outputs/historical_test_overlap_roster.csv');expected=s.read('final_evidence/benchmark_audit.csv').set_index('stage');rc=lambda x:x.translate(str.maketrans('ACGT','TGCA'))[::-1];rows=[]
for stage,m in zip([1,2],s.manifests()):
    test=roster[roster.stage.eq(stage)];train=pd.concat([m[~m.role.eq('legacy_test')],m[m.record_id.isin(test[test.exact_training_match].record_id)]])
    exact=test.sequence.isin(set(train.sequence));reverse=test.sequence.map(rc).isin(set(train.sequence));candidate=m[m.role.eq('train_candidate')]
    for seq,digest in zip(m.sequence,m.sequence_sha256):
        import hashlib
        assert hashlib.sha256(seq.encode()).hexdigest()==digest
    counts=dict(stage=stage,train=len(train),test=len(test),exact=int(exact.sum()),reverse_complement=int(reverse.sum()),candidates=len(candidate),groups=candidate.similarity_group_id.nunique())
    for a,b in [('train','supplied_train_n'),('test','supplied_test_n'),('exact','exact_train_test_overlap'),('reverse_complement','reverse_complement_train_test_overlap'),('candidates','candidate_n'),('groups','candidate_groups')]:assert counts[a]==expected.loc[stage,b]
    rows.append(counts)
s.dump('audit_validation.json',dict(passed=True,counts=rows,full_sequence_hashes_checked=True))
print('Benchmark counts and DNA hashes reproduced.')
'''

GROUPING='''import sys,itertools,hashlib
from pathlib import Path
import numpy as np
import pandas as pd
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from analysis_tools import final_synthesis as s
def lcs(a,b):
    masks={c:0 for c in 'ACGT'}
    for i,c in enumerate(a):masks[c]|=1<<i
    state=0
    for c in b:
        x=masks[c]|state;state=x&~(x-((state<<1)|1))
    return state.bit_count()
def rc(x):return x.translate(str.maketrans('ACGT','TGCA'))[::-1]
s.verify_inputs();rows=[]
for stage,m in zip([1,2],s.manifests()):
    candidate=m[m.role.eq('train_candidate')];assert candidate.groupby('similarity_group_id').fold_5.nunique().max()==1
    canonical=candidate.sequence.map(lambda x:min(x,rc(x)));assert candidate.assign(canonical=canonical).groupby('canonical').fold_5.nunique().max()==1
    checked=0
    if '--exhaustive' in sys.argv:
        # Recompute every pair among retained candidates. CPU-intensive; no training.
        records=list(candidate.itertuples())
        for a,b in itertools.combinations(records,2):
            if max(lcs(a.sequence,b.sequence),lcs(a.sequence,rc(b.sequence)))>=160:
                assert a.similarity_group_id==b.similarity_group_id
            checked+=1
    rows.append(dict(stage=stage,candidate_records=len(candidate),groups=candidate.similarity_group_id.nunique(),fold_isolation_passed=True,exhaustive_pairs_checked=checked))
# Independently recompute the archived positive Stage1 LCS edge values.
m=s.manifests()[0].drop_duplicates('record_id').set_index('record_id');edges=pd.read_csv(s.PACKAGE/'inputs/final_evidence/inputs/stage1_lcs_ge150_recomputed.tsv',sep='\\t');edge_checks=0
for row in edges.itertuples():
    if row.id_a not in m.index or row.id_b not in m.index:continue
    a=m.loc[row.id_a,'sequence'];b=m.loc[row.id_b,'sequence'];assert lcs(a,b)==row.lcs_forward and lcs(a,rc(b))==row.lcs_rc;edge_checks+=1
s.dump('grouping_validation.json',dict(passed=True,tasks=rows,positive_edges_recomputed=edge_checks,full_pair_rescan=('--exhaustive' in sys.argv),limitation='Default test verifies fixed-group fold isolation, exact/RC pairs and supplied positive edges; absence of additional edges requires --exhaustive. This does not rule out every biological relationship.'))
print('Fixed grouping isolation and positive LCS edge checks passed.')
'''

FIGURES='''import sys,json
from pathlib import Path
from PIL import Image,ImageChops
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from analysis_tools import final_synthesis as s
from analysis_tools import final_presentation as p
s.verify_inputs();p.phase43();rows=[]
for r in s.read('final_figures/figure_index.csv').itertuples():
    a=Image.open(s.PACKAGE/'inputs'/r.png).convert('RGB');b=Image.open(s.ROOT/r.png).convert('RGB');assert a.size==b.size
    assert ImageChops.difference(a,b).getbbox() is None,r.figure
    rows.append(dict(figure=r.figure,pixel_identical=True))
s.dump('figure_validation.json',dict(passed=True,figures=len(rows),checks=rows,meaning='All 13 PNG pixel arrays exactly reproduce in the recorded environment. PDF/SVG metadata bytes can vary.'))
print('All 13 scientific figures reproduced pixel-identically.')
'''

def phase47():
    s.verify_inputs();assert s.jd('phase46_outputs/summary.json')['status']=='completed_scientific_draft'
    RELEASE.mkdir(parents=True,exist_ok=True)
    for name in ['scripts','analysis_tools','data_manifests','predictions','tables','figures','docs','inputs']: (RELEASE/name).mkdir(exist_ok=True)
    dependencies=set()
    source=(ROOT/'analysis_tools/final_presentation.py').read_text(encoding='utf-8')
    tree=ast.parse(source);phase43=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='phase43')
    for node in ast.walk(phase43):
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Attribute) and node.func.attr in ['read','jd'] and node.args and isinstance(node.args[0],ast.Constant) and isinstance(node.args[0].value,str):dependencies.add(node.args[0].value)
    dependencies.update(['final_statistics/classification_predictions.csv.gz','phase37_outputs/verified_ensemble_predictions.csv.gz','final_statistics/classification_summary.csv','final_statistics/mpra_internal_summary.csv','final_statistics/external_transfer_summary.csv','final_statistics/statistical_methods.json','final_evidence/benchmark_audit.csv','final_evidence/inputs/stage1_candidate_manifest.csv','final_evidence/inputs/stage2_candidate_manifest.csv','final_evidence/inputs/stage1_lcs_ge150_recomputed.tsv','phase42_outputs/summary.json','final_figures/figure_index.csv'])
    dependencies.update(p.relative_to(ROOT).as_posix() for p in (ROOT/'phase35_outputs/training_logs').glob('*.csv'))
    dependencies.update(p.relative_to(ROOT).as_posix() for p in (ROOT/'final_figures').rglob('*.png') if p.name!='figure_contact_sheet.png')
    for relative in sorted(dependencies):copy_file(relative,RELEASE/'inputs'/relative)
    for p in (ROOT/'final_tables').rglob('*.csv'):copy_file(p.relative_to(ROOT),RELEASE/'tables'/p.relative_to(ROOT/'final_tables'))
    for p in (ROOT/'final_figures').rglob('*'):
        if p.is_file() and p.suffix in ['.png','.svg','.pdf','.md','.csv']:copy_file(p.relative_to(ROOT),RELEASE/'figures'/p.relative_to(ROOT/'final_figures'))
    for r,name in [('final_statistics/classification_predictions.csv.gz','classification_selected_ensembles_and_archived_baselines.csv.gz'),('phase37_outputs/verified_ensemble_predictions.csv.gz','MPRA_current_ensembles.csv.gz'),('phase32_outputs/independently_verified_predictions.csv.gz','classification_current_all_seeds.csv.gz')]:copy_file(r,RELEASE/'predictions'/name)
    for relative in ['phase30_outputs/training_protocol.json','phase34_outputs/training_protocol.json','phase35_outputs/run_lock.json','phase36_outputs/selection_lock.json','phase38_outputs/ERRATUM_collapse_count.md','phase32_outputs/checkpoint_reproduction_checks.csv','phase37_outputs/checkpoint_reproduction_checks.csv','phase30_outputs/pinned_pretrained_model_files.csv','phase34_outputs/copied_input_manifest.csv']:
        copy_file(relative,RELEASE/'docs/provenance'/Path(relative).parent.name/Path(relative).name)
    # Provenance snapshots are distinct from the portable reporting adapter.
    for name in ['mpra_fairness_training.py','regression_architectures.py','classification_repeated_training.py','final_synthesis.py','final_presentation.py']:
        p=ROOT/'analysis_tools'/name
        if p.exists():copy_file(p.relative_to(ROOT),RELEASE/'docs/source_snapshots'/name)
    write(RELEASE/'analysis_tools/__init__.py','')
    write(RELEASE/'analysis_tools/final_synthesis.py',ADAPTER)
    copy_file('analysis_tools/final_presentation.py',RELEASE/'analysis_tools/final_presentation.py')
    original=(ROOT/'analysis_tools/final_synthesis.py').read_text(encoding='utf-8');nodes=ast.parse(original)
    names=['classification_metrics','regression_metrics','weighted_corr','weighted_ranks','weighted_auc','weighted_metric','stats_selfcheck','paired_arrays']
    functions=['import numpy as np\nimport pandas as pd']+[ast.get_source_segment(original,n) for n in nodes.body if isinstance(n,ast.FunctionDef) and n.name in names]
    write(RELEASE/'analysis_tools/metric_functions.py','\n\n'.join(functions))
    for name,body in [('metrics.py',METRICS),('audit.py',AUDIT),('grouping.py',GROUPING),('figures.py',FIGURES)]:write(RELEASE/'scripts'/name,body)
    candidates=pd.concat(s.manifests(),ignore_index=True)
    s.table(RELEASE/'data_manifests/legacy_hashes.csv',candidates[['stage','record_id','sequence_sha256','original_split','role']])
    s.table(RELEASE/'data_manifests/grouped_folds.csv',candidates[candidates.role.eq('train_candidate')][['stage','record_id','sequence_sha256','label','similarity_group_id','fold_5']])
    copy_file('phase21_outputs/historical_test_overlap_roster.csv',RELEASE/'data_manifests/duplicate_manifest.csv')
    splits=s.read('phase34_inputs/phase9_frozen_splits.csv');s.table(RELEASE/'data_manifests/mpra_splits.csv',splits[splits.cell_line.isin(['HepG2','K562'])])
    versions={name:metadata.version(name) for name in ['numpy','pandas','scipy','scikit-learn','matplotlib','Pillow']}
    write(RELEASE/'requirements.txt','\n'.join(k+'=='+v for k,v in versions.items()))
    write(RELEASE/'environment.yml','name: enhancer\nchannels:\n  - conda-forge\ndependencies:\n  - python=3.12\n  - pip\n  - pip:\n'+''.join('      - '+k+'=='+v+'\n' for k,v in versions.items()))
    write(RELEASE/'LICENSE','LOCAL REVIEW CANDIDATE — LICENSE NOT YET GRANTED\n\nThe authors must confirm ownership and choose a license before public release. A permissive license such as MIT is proposed only for original code they own. Third-party architectures, pretrained model material, datasets and archived code retain their original terms. Inclusion here does not relicense them. No license permission or public-deposition approval is inferred from preparation of this candidate.')
    write(RELEASE/'docs/THIRD_PARTY_RIGHTS.md','''# Rights inventory and release decision

This is a local candidate for GitHub and an archival repository such as Zenodo. Neither repository nor DOI exists for this package yet. Do not insert an invented identifier into the manuscript.

| Component | Provenance | Release condition |
|---|---|---|
| Original synthesis/reporting code and derived results | Source snapshots, MANIFEST and claim ledger | Authors confirm ownership, contributors and license |
| Legacy benchmark sequences and labels | Supplied iEnhancer benchmark files; hashes and grouping manifests | Confirm redistribution terms and cite original source |
| Reporter activity and DHS-natural derived predictions | Frozen study cohorts; source/input manifests | Confirm original data permissions and cite assay source; predictions include observed activities |
| DNABERT / DNABERT-2 | Pinned revisions and pretrained file hashes | Retain upstream terms; no pretrained or fitted weights bundled |
| LegNet / MPRAnn architecture snapshots | Source URLs/commits in copied input manifest | Verify upstream code licenses and retain notices before public release |
| Fonts and Python libraries | Environment file | Install separately under their licenses; no runtime binaries bundled |

Feedback correspondence, reviewer letters, author email addresses, large weights, raw workspace caches and previous sealed archives are excluded. Data-derived DNA sequences and activity measurements are present in this local candidate and need rights review before public upload. No global installation is performed by this workflow.
''')
    write(RELEASE/'README.md','''# Enhancer benchmark dependence and training adequacy

Local release candidate for the manuscript “Benchmark dependence and training adequacy shape enhancer sequence model evaluation”. This compact package reconstructs fixed-prediction metrics, final conditional statistics, benchmark count/hash audits, fixed-group isolation checks, and all six main plus seven supplementary figures. It does not retrain models or reproduce checkpoint inference without separately archived weights and full training inputs.

## Run

Use Python 3.12 with requirements.txt (tested versions). From this package directory:

```powershell
python -B -X utf8 scripts/audit.py
python -B -X utf8 scripts/grouping.py
python -B -X utf8 scripts/metrics.py --full-statistics
python -B -X utf8 scripts/figures.py
```

Outputs stay in `reproduced/`; input hashes are checked first. The statistics command reruns all 14 contrasts with 2,000 group-bootstrap and 9,999 paired prediction-swap draws each, then Holm adjustment. It can take several minutes. `grouping.py --exhaustive` additionally rescans every retained sequence pair using exact LCS in both orientations; the default validation does not claim a fresh exhaustive rescan.

The reproducibility report records the actual local tests. Prediction tables distinguish current three-seed ensembles from archived single runs. Tables retain all failed/near-constant fits and the original eligibility policy. Zero original undefined-Pearson flags and two practical near-constant fits are different counts.

## Structure

`data_manifests/`: identities, duplicate roster, fixed folds and MPRA splits. `scripts/`: executable checks. `predictions/`: compressed scored records. `tables/` and `figures/`: expected scientific outputs. `inputs/`: exact small reporting inputs mirrored by their source-relative names. `docs/`: reproduction limits, protocol/checkpoint hashes, rights and source snapshots.

## Interpretation

All cohorts were previously examined. Intervals and paired tests are conditional on fixed fitted models; they exclude model-selection and training-algorithm uncertainty. Folds and seed-by-fold points are dependent. Equal candidate counts are unequal compute. Pointwise intervals and family-adjusted tests are distinct; a nonsignificant contrast does not establish equivalence. Ranking does not establish absolute calibration or endogenous regulatory function.

## Public release

This local candidate has not been uploaded. Author approval, third-party rights review and a final code/data license are required. See LICENSE and docs/THIRD_PARTY_RIGHTS.md. No public URL or DOI is claimed.
''')
    copy_file('final_supplement/Supplementary_Methods.md',RELEASE/'docs/reproducibility.md')
    copy_file('claim_audit/claim_evidence_ledger.csv',RELEASE/'docs/claim_evidence_ledger.csv')
    write(RELEASE/'docs/RETRAINING_AND_CHECKPOINTS.md','''# Scope of reproducibility

The compact commands have been tested for reporting and fixed-prediction reproduction. Source snapshots of the frozen implementation, pinned source/model metadata, protocols, split/sequence hashes and checkpoint reproduction indices are supplied. Fitted and pretrained weights, optimizer/resume states and full raw training datasets remain in the preserved internal study. The original checks reproduced 60 classification checkpoints and 72 MPRA validation checkpoints plus 60 evaluation sets; this compact release does not repeat that GPU inference or claim to bundle those weights.

Full training requires separately obtaining authorized data, the exact pinned pretrained files and original study layout, then using the frozen training source and protocols. The compact package is not a standalone end-to-end retraining distribution. Publicly archive separately licensed checkpoints/data, or document their access conditions, before claiming a complete training release. Do not run training source against an approximate replacement layout or substitute unavailable historical implementations. The historical three-model leakage-impact experiment remains stopped for unavailable original neural provenance.
''')
    write(RELEASE/'MANIFEST.json',json.dumps(manifest(RELEASE),indent=2))
    env=os.environ.copy();env.update(PYTHONDONTWRITEBYTECODE='1',PYTHONUTF8='1',MPLCONFIGDIR=str(RELEASE/'.runtime/matplotlib'),TEMP=str(RELEASE/'.runtime/temp'),TMP=str(RELEASE/'.runtime/temp'))
    for p in [RELEASE/'.runtime/matplotlib',RELEASE/'.runtime/temp']:p.mkdir(parents=True,exist_ok=True)
    tests=[]
    for name,args in [('audit.py',[]),('grouping.py',[]),('metrics.py',['--full-statistics']),('figures.py',[])]:
        log=ROOT/f'phase47_outputs/{name}.log';log.parent.mkdir(exist_ok=True)
        print('Portable test:',name,flush=True)
        with log.open('w',encoding='utf-8') as stream:
            result=subprocess.run([sys.executable,'-B','-X','utf8',str(RELEASE/'scripts'/name),*args],cwd=ROOT,env=env,stdout=stream,stderr=subprocess.STDOUT)
        assert result.returncode==0,log
        tests.append(dict(script=name,args=args,exit_code=0,log=log.relative_to(ROOT).as_posix()))
    validations={name:json.loads((RELEASE/'reproduced'/name).read_text()) for name in ['audit_validation.json','grouping_validation.json','metric_validation.json','figure_validation.json']}
    assert all(v['passed'] for v in validations.values())
    report=dict(passed=True,tests=tests,validation=validations,versions=versions,python=sys.version,working_directory_for_local_test=str(ROOT),all_outputs_within_contribution=True,new_training=False,public_upload=False)
    # Public candidate omits the local absolute test cwd.
    public_report={**report,'working_directory_for_local_test':'contribution root; portable scripts resolve inputs from their own location'}
    write(RELEASE/'docs/REPRODUCIBILITY_TEST.json',json.dumps(public_report,indent=2))
    s.dump('phase47_outputs/reproducibility_validation.json',report)
    package=seal(RELEASE,ROOT/'phase47_outputs/Enhancer_Reproducibility_Release_Candidate.zip')
    s.dump('phase47_outputs/summary.json',dict(status='completed_local_release_candidate',publicly_uploaded=False,license_and_redistribution_approval_pending=True,full_training_weights_bundled=False,portable_fixed_prediction_tests_passed=True,package=package,source_sha256=s.sha(__file__)))
    print('Phase47 local release candidate completed.',flush=True)

def phase48():
    """A local journal-neutral review bundle; author assertions remain unmade."""
    s.verify_inputs();assert s.jd('phase47_outputs/summary.json')['portable_fixed_prediction_tests_passed']
    assert s.jd('phase46_outputs/visual_QA.json')['passed']
    from analysis_tools.final_documents import runtime
    runtime()
    folder=ROOT/'submission_package';folder.mkdir(exist_ok=True)
    for source,target in [('manuscript/Manuscript.docx','Manuscript.docx'),('manuscript/Manuscript.pdf','Manuscript.pdf'),('manuscript/Manuscript.md','Manuscript.md'),('final_supplement/Supplementary_Information.pdf','Supplementary_Information.pdf'),('final_supplement/Supplementary_Tables.xlsx','Supplementary_Tables.xlsx')]:copy_file(source,folder/target)
    for source_dir,target_dir in [('final_figures/main','Main_Figures'),('final_figures/supplement','Supplementary_Figures'),('final_tables/main','Main_Tables')]:
        for p in (ROOT/source_dir).iterdir():
            if p.is_file():copy_file(p.relative_to(ROOT),folder/target_dir/p.name)
    for source in ['final_figures/FIGURE_CAPTIONS.md','final_supplement/Supplementary_Methods.md','final_supplement/Supplementary_Result_Notes.md','manuscript/references.csv','claim_audit/claim_evidence_ledger.csv','claim_audit/actual_manuscript_paragraph_checks.csv']:
        copy_file(source,folder/'Evidence_and_Methods'/Path(source).name)
    # Complete the release provenance with the actual frozen classification source.
    class_lock=s.jd('phase31_outputs/run_lock.json');matches=[]
    for p in (ROOT/'analysis_tools').glob('*.py'):
        if s.sha(p)==class_lock['source_sha256']:matches.append(p)
    assert len(matches)==1,'Frozen classification source must be identified by hash.'
    copy_file(matches[0].relative_to(ROOT),RELEASE/'docs/source_snapshots'/matches[0].name)
    copy_file('phase31_outputs/run_lock.json',RELEASE/'docs/provenance/phase31_outputs/run_lock.json')
    old=s.jd('phase47_outputs/summary.json')
    old['classification_source_snapshot_added_by_hash']=dict(path=matches[0].relative_to(ROOT).as_posix(),sha256=class_lock['source_sha256'])
    old['package']=seal(RELEASE,ROOT/'phase47_outputs/Enhancer_Reproducibility_Release_Candidate.zip')
    s.dump('phase47_outputs/summary.json',old)
    copy_file('phase47_outputs/Enhancer_Reproducibility_Release_Candidate.zip',folder/'Reproducibility/Enhancer_Reproducibility_Release_Candidate.zip')
    write(folder/'Highlights.md','''# Highlights

- Supplied benchmark test files contain 100 training-identical strong records.
- Grouping and exact-GC evaluation change the interpretation of classification scores.
- Longer DNABERT-2 training improves its own short-prefix MPRA predictions.
- LegNet remains stronger under the tested validation-selected internal recipes.
- External ranking persists despite negative R² and systematic underprediction.
''')
    write(folder/'Contribution_Box.md','''# Contribution statement for manuscript planning

The study contributes an integrated, traceable evaluation of benchmark nonindependence, composition-sensitive classification, repeated transformer training, bounded activity-model searches, own-prefix duration controls and assay transfer. Its main value is the distinction between questions that a single leaderboard conflates: benchmark validity, quantitative prediction, training adequacy and calibration under transfer.

The study does not introduce a new architecture or establish first priority for a language-model/CNN comparison on MPRA data. Equal candidate counts do not imply equal compute or exhaustive tuning. The preserved historical leakage-impact experiment stops where original neural provenance is unavailable. Current repeated classification, current MPRA fits and older filtering/reliability diagnostics remain separate evidence generations.
''')
    write(folder/'Data_and_Code_Availability_DRAFT.md','''# Data and Code Availability — local draft

The analyses use supplied legacy enhancer classification files and previously generated reporter-activity and DHS-natural cohorts. Sequence identities, original source metadata, duplicate rosters, fixed group/fold assignments, MPRA split manifests, processed predictions and SHA256 provenance are organized in the accompanying local reproducibility release candidate. The supplementary methods identify assay sources and distinguish raw/retained populations, historical fitted models and current experiments.

Analysis scripts, tested environment specifications, final tables/figures, fixed-prediction outputs and frozen source/protocol snapshots are included in `Reproducibility/Enhancer_Reproducibility_Release_Candidate.zip`. That candidate has been tested for metric/statistical reproduction, benchmark auditing and scientific figure regeneration. It excludes pretrained/fitted weights and does not claim self-contained retraining or fresh GPU checkpoint inference. Separate checkpoint hash indices document the previously completed checkpoint audits.

Public repository and archival deposition are pending author permission, redistribution checks and a final license. No public repository URL, accession or DOI has yet been assigned. Before submission, replace this local-status description with verified public links or accurate access conditions and clearly disclose any unavailable training inputs/checkpoints. Do not state that all materials are publicly available while they remain only on this computer.
''')
    write(folder/'Author_Declarations_REQUIRED.md','''# Author confirmations required before submission

The scientific draft is prepared locally. The following factual declarations cannot be inferred from the analysis or silently copied from an earlier manuscript.

1. Final author names, order, affiliations, ORCIDs, corresponding author and contact details. Confirm Shafin's full name, eligibility for authorship and role. An earlier draft lists Arian Rahman Aditta, Udita Rahman Arshee, Sheikh Md. Rabiul Islam and Nishat Taaha; that is historical metadata, not confirmation of the final list. One historical email address appears incomplete and needs correction.
2. CRediT roles for each final author: conceptualization; methodology; software; validation; formal analysis; investigation; resources; data curation; writing — original draft; writing — review and editing; visualization; supervision; project administration; funding acquisition. Assign only roles each author confirms. Saved code or an assistant's work log is not evidence of every person's contribution.
3. Funding agency, grant numbers, institutional support and funder role, or an explicit author-confirmed no-specific-funding statement.
4. Financial and nonfinancial competing interests, or an author-confirmed no-competing-interests statement.
5. Approval of the manuscript by all authors and accountability for its content. No approval is asserted by this draft.
6. Prior dissemination, preprint status, related manuscripts and current submission status at Array. Do not assert exclusive submission, withdrawal or previous rejection without confirmation. The Array response matrix is an internal conditional draft.
7. Data/model/code redistribution permissions, repository owners, final licenses and the intended archival deposit. No upload or license grant has been made.
8. Any ethics/consent requirement or exemption applicable to the original sources. This analysis used existing sequence and reporter datasets; do not invent an institutional approval number.
9. Generative-AI disclosure: confirm the actual tools, provider/version where known, dates and uses; verify all scientific claims/citations/code and describe human oversight. Proposed wording is supplied below for author review.

## Proposed AI-use disclosure for review

“An OpenAI coding/writing assistant was used to assist with data auditing, analysis and reporting code, figure/table preparation, and drafting or editing manuscript text. The authors reviewed the source evidence, statistical methods, outputs and citations and take responsibility for the final work.”

The second sentence must only be used after the authors actually complete that review. An AI assistant is not an author. Confirm any additional AI tools used earlier in the project; this draft does not assert a complete historical tool inventory.

## CRediT worksheet

Prepare one row per confirmed author with their approved roles. Authorship, order and roles remain pending; no person has been assigned a role by inference in this package.
''')
    credit=pd.DataFrame([dict(author_name=name,author_order='pending confirmation',affiliation='pending confirmation',ORCID='pending confirmation',CRediT_roles='pending author confirmation',approved_final_manuscript='not yet confirmed') for name in ['Arian Rahman Aditta','Udita Rahman Arshee','Sheikh Md. Rabiul Islam','Nishat Taaha','Shafin — full name and authorship decision required']])
    s.table(folder/'CRediT_Author_Confirmation_Worksheet.csv',credit)
    write(folder/'Target_Journal_Decision.md','''# Target-journal decision

Scientific Reports is the practical recommendation for this study's present scope: benchmark audit, reproducible comparisons, training adequacy and bounded assay-transfer findings. Its publisher emphasizes technical validity and supported conclusions across scientific disciplines. PLOS Computational Biology is a more ambitious option: its guidelines accept benchmarking, but expect substantial new computational/biological insight. No acceptance probability is implied.

The Scientific Reports publisher information sheet inspected for this recommendation identifies Q1 in Multidisciplinary Sciences in the **2025 Journal Citation Reports**. That supports a dated Q1 statement, not an independently verified claim about the newest JCR edition or every subject category. Quartiles must name the database, category and year.

Sources checked 15 October 2026:
- https://www.nature.com/srep/about
- https://www.nature.com/documents/063F5_Infographc_297x210mm_proof5.pdf
- https://journals.plos.org/ploscompbiol/s/submission-guidelines
- https://journals.plos.org/ploscompbiol/s/journal-information

The user has not yet selected the final journal. This bundle is journal-neutral. After that choice, check that journal's current article type, author/declaration instructions, reference style, figure/file requirements, charges and publication options before creating a submission-ready version. No submission, journal contact or payment is performed here.
''')
    cover='''Dear Editor,

Please consider “Benchmark dependence and training adequacy shape enhancer sequence model evaluation” for publication as a research article.

The manuscript integrates a sequence-level benchmark audit with similarity-grouped classification, exact-GC evaluation, repeated DNABERT comparisons and a bounded validation-only search for reporter-activity prediction. It identifies 100 training-identical strong test records in the supplied legacy benchmark files, and separates this reproducible overlap finding from the causal effect that cannot be recovered for unavailable historical neural implementations.

A 54-fit MPRA study compares DNABERT-2, LegNet and MPRAnn using equal candidate/seed counts, with an own-prefix duration control. Longer DNABERT-2 training improves internal Pearson correlation by 0.0534 in HepG2 and 0.0501 in K562 over its exact early prefixes, while LegNet remains stronger under the tested selected recipes. External rank preservation coexists with negative R², illustrating why discrimination, training adequacy and absolute transfer should be evaluated separately.

The contribution is an integrated evaluation and provenance framework, rather than a new architecture or a universal model ranking. The paper discloses the previously examined cohorts, conditional uncertainty, dependent folds, limited search and practical near-constant fits. A methods supplement, processed predictions, tables, figures and locally tested reproduction scripts accompany the draft.

The author team will complete the required authorship, funding, competing-interest, prior-submission and data/code-availability declarations before submission. This local letter draft does not assert author approval, exclusivity or public deposition.
'''
    write(folder/'Cover_Letter_DRAFT.md','# Cover letter — unsigned local draft\n\n'+cover)
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    document=Document();sec=document.sections[0];sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
    normal=document.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(8)
    document.add_heading('Cover letter — unsigned draft',0)
    for para in cover.strip().split('\n\n'):document.add_paragraph(para)
    document.core_properties.author='';document.core_properties.last_modified_by='';document.save(folder/'Cover_Letter_DRAFT.docx')
    from analysis_tools.final_manuscript import render_docx
    letter_pages=render_docx('submission_package/Cover_Letter_DRAFT.docx','phase48_outputs/rendered_cover_letter')
    copy_file('phase48_outputs/rendered_cover_letter/Cover_Letter_DRAFT.pdf',folder/'Cover_Letter_DRAFT.pdf')
    # Scientific graphical abstract, derived directly from locked contrast estimates.
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    fig,ax=plt.subplots(figsize=(12,5.2));ax.set(xlim=(0,12),ylim=(0,5.2));ax.axis('off')
    ax.text(6,4.92,'Benchmark validity, training adequacy and assay transfer',ha='center',va='center',fontsize=16,fontweight='bold')
    panels=[(.2,'Audit and group','100 strong test records\nidentical to training\n\nFixed similarity groups\nand exact-GC evaluation','#275D8C'),(4.2,'Separate training questions','Longer DNABERT-2: internal r\n+0.0534 HepG2 / +0.0501 K562\n\nSelected LegNet remains stronger\nunder the tested recipes','#35806A'),(8.2,'Evaluate transfer honestly','External rank correlations > 0\nbut all selected external R² < 0\n\nRanking does not establish\nabsolute calibration','#C65A28')]
    for x,title,body,color in panels:
        ax.add_patch(FancyBboxPatch((x,1.2),3.6,3.05,boxstyle='round,pad=0.06',facecolor='#F4F6F8',edgecolor=color,lw=1.4))
        ax.text(x+1.8,3.85,title,ha='center',va='center',fontsize=12,fontweight='bold',color=color)
        ax.text(x+1.8,2.62,body,ha='center',va='center',fontsize=10.5,linespacing=1.7)
    for x in [3.84,7.84]:ax.annotate('',xy=(x+.31,2.72),xytext=(x,2.72),arrowprops=dict(arrowstyle='->',lw=1.7,color='#59636D'))
    ax.text(6,.62,'Previously examined cohorts · fixed-model conditional inference\nEqual candidate counts, unequal compute · no universal architecture ranking',ha='center',va='center',fontsize=10,color='#485566',linespacing=1.6)
    for extension in ['png','svg','pdf']:fig.savefig(folder/f'Graphical_Abstract.{extension}',dpi=220,bbox_inches='tight',facecolor='white')
    plt.close(fig)
    write(folder/'Graphical_Abstract_Caption.md','The diagram summarizes separate evidence questions. Internal duration changes are Pearson-correlation differences for fixed-learning-rate longer DNABERT-2 versus its own exact early-prefix ensembles. LegNet comparisons follow validation-only selection. All selected external R² values are negative. The diagram does not estimate causal leakage inflation or universal architecture superiority.')
    changes=[
      ('Benchmark independence','Added exact/RC audit, source-specific counts, retained graph groups and honest historical stop rule.','Main Tables 1–2; Figure 2; Supplementary Methods; phase40 STOP_RULE_AND_LIMITS.'),
      ('Classification and composition','Current repeated DNABERT fits; archived baselines separate; exact-GC matching; calibration/confusion and mixed-label sensitivity.','Main Table 3; Figures 3–4; Supplementary Tables S2–S5 and Figures SF1–SF2.'),
      ('Training adequacy and model comparison','54 validation-only fits, own-prefix duration controls, selected epoch/grid/cap diagnostics, collapse erratum and recorded cost.','Main Table 4; Figure 5; Supplementary Tables S6–S8,S16–S17; Figures SF3–SF5,SF7.'),
      ('Transfer and reliability','Positive rank correlation versus negative external R²; older filtering/tail/coverage analyses explicitly labelled.','Main Table 5; Figure 6; Supplementary Tables S9–S14; Figure SF6.'),
      ('Statistics and uncertainty','Fourteen final-synthesis contrasts with conditional group intervals, whole-group swaps and Holm adjustment; no equivalence claim.','Main Table 6; Supplementary Table S18; statistical_methods.json.'),
      ('Claims and prior literature','80 verified numerical claim records; current evidence hierarchy; bounded contribution; 18 references.','Claim ledger; actual manuscript checks; manuscript Discussion and References.'),
      ('Reproducibility','Compact local release with executed metrics/statistics/audit/figure reproduction; no public URL or bundled weights claimed.','Reproducibility release candidate and test report.'),
      ('Submission facts','Prepared neutral local cover letter, highlights, graphical abstract, declarations worksheet and reporting map.','Author and journal decisions remain pending; no submission performed.')]
    changes_frame=pd.DataFrame(changes,columns=['topic','completed_change','evidence_location']);s.table(folder/'Summary_of_Changes.csv',changes_frame)
    write(folder/'Summary_of_Changes.md','# Summary of scientific and reporting changes\n\n'+'\n\n'.join('**'+a+'**\n\n'+b+'\n\nEvidence: '+c for a,b,c in changes)+'\n\nThe optional historical three-model experiment was assessed and stopped, not executed, because the original neural historical implementation/checkpoints/predictions were unavailable. No new architectures, long-context inputs, cell lines or arbitrary seeds were added.')
    reporting=[('DOME: data','Identity hashes, inclusion/exclusion, folds, units and label composition','Methods; Tables 1–2; Supplementary Methods and S1,S3,S10','Covered within stated scope; original publisher provenance discrepancy unresolved'),('DOME: optimization','Fixed classification protocol and validation-only MPRA grids/seed/checkpoint policies','Methods; S6–S8,S16–S17; protocol snapshots','Limited candidates/compute; selected budget/grid flags retained'),('DOME: model','Pinned revisions, architectures, source and checkpoint hashes','Supplementary Methods; provenance/source snapshots','Weights not bundled in compact candidate'),('DOME: evaluation','Fixed cohorts, independent metric checks, conditional group inference and failure diagnostics','Tables 3–6; S2–S5,S15,S18','Previously examined cohorts; no fresh confirmation'),('REFORMS: leakage and validity','Exact/RC/LCS grouping audit; train-only normalization and pre-inference recipe locks','Methods; benchmark audit; preserved locks','Does not exclude every biological relation or pretraining overlap'),('REFORMS: uncertainty and reproducibility','Separate seed/fold/conditional summaries; full portable reporting tests','Supplement; local release validation','No algorithm/search uncertainty or universal claims'),('REFORMS: transparency and availability','Historical stop rule, collapse erratum, code/data/rights inventories','Results; Data/Code draft; release docs','Public deposition, licenses and author approval pending')]
    s.table(folder/'DOME_REFORMS_Reporting_Map.csv',pd.DataFrame(reporting,columns=['reporting_topic','implementation','location','limit']))
    write(folder/'DOME_REFORMS_Reporting_Map.md','# Reporting map\n\nThis is a topic map, not a claim of checklist certification or complete elimination of leakage.\n\n'+'\n\n'.join('**'+a+'**: '+b+'. Location: '+c+'. Limit: '+d+'.' for a,b,c,d in reporting))
    write(folder/'Array_Response_DRAFT_NOT_FOR_SENDING.md','''# Conditional Array response — internal drafting aid

Use only if the authors choose to return to Array and confirm the editorial status. No response has been sent. This is not a fabricated verbatim reviewer response: the rows summarize scientific issues in the preserved review/evidence work and must be aligned with the exact reviewer numbering before sending.

| Issue | Evidence-backed response | Remaining limit |
|---|---|---|
| Independent benchmark evaluation | Exact/RC identity audit and fixed LCS component grouping with traceable exclusions | Original neural historical implementation unavailable; no causal inflation estimate |
| Composition confounding | Exact-GC pair AUROC, archived composition baselines, eligibility counts | Evaluation sensitivity, not causal control of every feature |
| Repeated model comparisons | 60 current consistent classification fits and separate seed/fold summaries | Three seeds; dependent folds; no equivalence |
| Reporter activity modelling | Selected DNABERT-2, LegNet and MPRAnn, frozen chromosome partitions, train-only scaling | Two cells, short sequences and finite recipes |
| Too-short DNABERT training | Exact own-prefix duration controls show improvements in both cells | Selected cap/grid flags prevent convergence claims |
| Reliability and transfer | Negative external R² and older tail/filtering/coverage failures retained | Older diagnostics are not new selected-model inference |
| Reproducibility and claims | Checkpoint audits preserved; fixed-prediction reproduction tested; claims traced | Public deposit, rights and author declarations pending |

Any editorial statement about withdrawal, novelty relative to a prior submission, author approval or exclusivity requires factual author confirmation. Consult the original review document and the summary-of-changes file when drafting the final point-by-point response.
''')
    pending=['Final journal and article type','Final authors/order/corresponding contact and CRediT','Funding, conflicts, ethics applicability and AI-use verification','All-author approval and Array/exclusivity status','Code/data licenses and public repository/archive access','Journal-specific file/reference checks after journal choice']
    write(folder/'START_HERE.md','''# Local manuscript review package

Read `Manuscript.pdf` first: 16 pages of text, references, six logical tables and figure legends, followed by six full-size main figures (22 pages total). `Manuscript.docx` remains editable. The supplement contains 15 PDF pages, 18 data tables plus a contents sheet, and seven scientific figures in separate high-resolution/vector files.

`Summary_of_Changes.md` and `Contribution_Box.md` explain the paper's evidence and contribution. The 14 headline conditional contrasts, source hashes and actual manuscript claim links are traceable. `Reproducibility/` contains the compact local release candidate and its tested scripts.

This is a scientific draft and local review bundle, not a submitted manuscript. The target journal is not yet selected; Scientific Reports is the practical recommendation. Public deposition, licenses and author factual declarations remain pending. No external message or upload was made. The historical three-model leakage-impact experiment stopped under its authentication rule; it was not silently replaced.

Before submission, complete Author_Declarations_REQUIRED.md and the CRediT worksheet, select the journal, obtain all-author approval, and replace local-only availability wording with verified public access conditions. Keep the original contribution workspace and older sealed studies intact.
''')
    s.dump('phase48_outputs/summary.json',dict(status='assembled_pending_visual_QA_and_final_validation',journal_target='not selected',cover_letter_pages=letter_pages,main_figures=6,supplementary_figures=7,pending_author_or_journal_facts=pending,public_upload=False,submission=False,new_training=False,source_sha256=s.sha(__file__)))
    print('Phase48 assembled; inspect graphical abstract and cover-letter pages before sealing.',flush=True)

def finalize48():
    """Seal only after actual notebook execution and visual review."""
    s.verify_inputs();from analysis_tools.final_documents import runtime
    runtime();import nbformat
    from pypdf import PdfReader
    qa=s.jd('phase48_outputs/visual_QA.json');assert qa['passed']
    for phase in [43,44,46]:
        path='final_figures/visual_QA.json' if phase==43 else f'phase{phase}_outputs/visual_QA.json'
        assert s.jd(path)['passed']
    assert s.jd('phase47_outputs/reproducibility_validation.json')['passed']
    assert s.jd('phase44_outputs/workbook_value_validation.json')['passed']
    notebooks=[]
    for number in range(39,49):
        matches=list(ROOT.glob(f'[{chr(91)}]{number}]*.ipynb'))
        # Brackets are literal filename characters, not a notebook glob class.
        matches=[p for p in ROOT.glob('*.ipynb') if p.name.startswith(f'[{number}] ')]
        assert len(matches)==1,(number,matches)
        n=nbformat.read(matches[0],as_version=4);cells=[c for c in n.cells if c.cell_type=='code']
        assert cells and all(c.execution_count is not None for c in cells),number
        assert not any(o.output_type=='error' for c in cells for o in c.outputs),number
        notebooks.append(dict(phase=number,path=matches[0].relative_to(ROOT).as_posix(),sha256=s.sha(matches[0]),executed_code_cells=len(cells),errors=0))
    assert s.sha(ROOT/'phase34_outputs/training_protocol.json')=='2d204c4d565b46ebbb1833a00141d045201d59b79b503d43c3e6d9214d1101a9'
    assert s.sha(ROOT/'analysis_tools/mpra_fairness_training.py')=='83a1ae9a9e03361e4fe517044ee7b3cc76f580db04878e4a12f1f7873c60d311'
    assert len(PdfReader(ROOT/'manuscript/Manuscript.pdf').pages)==22
    assert len(PdfReader(ROOT/'final_supplement/Supplementary_Information.pdf').pages)==15
    ledger=s.read('claim_audit/claim_evidence_ledger.csv');assert ledger.status.eq('VERIFIED').sum()==80
    # No prior contribution entry is rewritten or silently removed.
    logpath=ROOT/'CONTRIBUTION_LOG.md';before=logpath.read_bytes();baseline=(ROOT/'final_evidence/contribution_log_at_start.md').read_bytes()
    assert before.startswith(baseline),'Contribution log differs from the preserved initial prefix.'
    marker='## Phases 39–48 — final evidence synthesis and local manuscript package (15 October 2026)'
    if marker.encode() not in before:
        entry=f'''\n\n{marker}

Completed the authorized final evidence freeze, statistical harmonization, six main/18 supplementary tables, six main/seven supplementary figures, supplementary workbook/PDF, prewriting and actual-manuscript claim audits, scientific manuscript, compact reproduction candidate and local submission-review bundle. All work remained inside Shafin's Contribution; earlier Phases12–38 and sealed studies were preserved.

The optional historical three-model leakage-impact experiment was assessed and stopped under its authentication rule: original neural historical implementations/checkpoints/independent-test predictions were unavailable. The separately authenticated archived spaced-mer predictions were rescored descriptively, with cleaned Stage2 discrimination undefined because no positives remain. No approximate neural reconstruction or causal inflation estimate is claimed.

Final synthesis uses 14 contrasts with 2,000 whole-group bootstrap and 9,999 conditional paired prediction-swap draws per contrast, plus Holm adjustment. The contrast family was fixed for final synthesis after prior outcomes were seen; all cohorts were previously examined. Portable reproduction reran the full contrast computations, 18 classification/20 MPRA ensemble metric rows, benchmark count/hash/group checks and all 13 scientific figures. The saved validation report states exact errors and pixel checks; this compact package does not repeat training/checkpoint inference or bundle model weights.

Current repeated classification DNABERT-2 minus DNABERT pooled AUROC differences are 0.0016 and 0.0125; both conditional 95% intervals include zero, without implying equivalence. Fixed-lr DNABERT-2 longer-versus-own-prefix internal Pearson differences are +0.0534 HepG2 and +0.0501 K562. Selected LegNet exceeds selected DNABERT-2 internally by +0.0389 and +0.0255 under the tested finite search. Selected external R² values are all negative despite positive rank correlations. The report retains zero frozen undefined-Pearson flags, two practical near-constant fits, six selected-fit budget flags and LR-grid boundary limits. Older filtering/reliability diagnostics remain separate fitted-model evidence.

The manuscript contains 31 evidence-linked paragraphs and 18 references. The editable DOCX and 22-page review PDF (16 manuscript/table/legend pages plus six full-size figures) were checked; the 15-page supplement and 19-sheet workbook were inspected, including independent saved-XLSX verification of 34,729 data/header cells. The ledger retains 80 VERIFIED numerical claim records and bounded/removed claims separately. Graphical abstract and unsigned cover-letter draft were visually checked before sealing.

Paths: manuscript/; final_evidence/; final_statistics/; final_tables/; final_figures/; final_supplement/; claim_audit/; public_release/enhancer-leakage-study/; submission_package/; executed notebooks39–48; phase47_outputs/Enhancer_Reproducibility_Release_Candidate.zip; phase48_outputs/Enhancer_Manuscript_Review_and_Submission_Candidate.zip. Final validation and archive hashes: phase48_outputs/FINAL_VALIDATION.json.

Status: scientific draft and local review packages completed; public upload and submission NOT performed. Final journal selection, author/order/CRediT/contact approval, funding/conflict/ethics/AI declarations, prior Array submission status, licenses/redistribution rights and public repository/archive identifiers remain author decisions. Scientific Reports was recommended with publisher-documented 2025-JCR Q1 context; no newest-edition or acceptance claim is made. Journal-specific tailoring remains pending the user's choice. No new model/cell/context training was performed in this synthesis.
'''
        with logpath.open('ab') as f:f.write(entry.encode('utf-8'))
    assert logpath.read_bytes().startswith(before)
    folder=ROOT/'submission_package';copy_file('CONTRIBUTION_LOG.md',folder/'Evidence_and_Methods/CONTRIBUTION_LOG.md')
    s.dump('phase48_outputs/notebook_execution_validation.json',notebooks)
    copy_file('phase48_outputs/notebook_execution_validation.json',folder/'Evidence_and_Methods/notebook_execution_validation.json')
    for relative in ['phase43_outputs/summary.json','phase44_outputs/visual_QA.json','phase44_outputs/workbook_value_validation.json','phase46_outputs/visual_QA.json','phase47_outputs/reproducibility_validation.json','phase48_outputs/visual_QA.json']:
        copy_file(relative,folder/'Evidence_and_Methods/validation'/Path(relative).parent.name/Path(relative).name)
    # Freeze the delivered reporting implementation and executed notebooks separately.
    snapshot=ROOT/'phase48_outputs/delivered_source';snapshot.mkdir(exist_ok=True)
    names=['final_synthesis.py','final_presentation.py','final_documents.py','final_claims.py','final_manuscript.py','final_delivery.py','validate_final_artifacts.py','build_final_supplement.mjs','export_final_docx_word.ps1']
    for name in names:copy_file(Path('analysis_tools')/name,snapshot/name)
    for n in notebooks:copy_file(n['path'],snapshot/Path(n['path']).name)
    source_manifest=manifest(snapshot);s.dump('phase48_outputs/delivered_source_manifest.json',source_manifest)
    package=seal(folder,ROOT/'phase48_outputs/Enhancer_Manuscript_Review_and_Submission_Candidate.zip')
    report=dict(status='completed_local_scientific_and_review_packages',scientific_checks_passed=True,all_notebooks39_to48_executed_without_errors=True,protected_history_passed=True,frozen_training_source_and_protocol_unchanged=True,main_tables=6,supplementary_tables=18,main_figures=6,supplementary_figures=7,manuscript_review_pdf_pages=22,supplement_pdf_pages=15,verified_numerical_claim_records=80,contribution_log_appended=True,contribution_log_sha256=s.sha(logpath),release_package=s.jd('phase47_outputs/summary.json')['package'],submission_review_package=package,public_upload=False,submission=False,journal_specific_tailoring_pending=True,author_declarations_and_licenses_pending=True,optional_historical_experiment='authentication stop rule applied; not run',new_training=False)
    s.dump('phase48_outputs/FINAL_VALIDATION.json',report)
    summary=s.jd('phase48_outputs/summary.json');summary.update(status='completed_local_review_package',visual_QA_passed=True,final_validation='phase48_outputs/FINAL_VALIDATION.json',package=package);s.dump('phase48_outputs/summary.json',summary)
    print('Final local review bundle validated and sealed; author/journal/public-release decisions remain pending.',flush=True)
